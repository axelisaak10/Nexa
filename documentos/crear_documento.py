from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.8)
sec.left_margin = sec.right_margin = Inches(.85)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title',26),('Heading 1',16),('Heading 2',12)]:
    st=doc.styles[name]
    st.font.name='Calibri'
    st.font.size=Pt(size)
    st.font.color.rgb=RGBColor(0,0,0)

def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text,1)
def sub(text): doc.add_heading(text,2)
def label(title,text):
    q=doc.add_paragraph()
    q.add_run(title+' ').bold=True
    q.add_run(text)

doc.add_paragraph('Actualización móvil de Nexa por etapas', 'Title')
h('Objetivo de la actualización')
p('La actualización propuesta de Nexa busca facilitar el acceso, la búsqueda de productos, las compras y el seguimiento de pedidos mediante las capacidades de los smartphones. Su implementación se organizará en seis etapas para validar cada función antes de incorporar la siguiente.')
h('Por qué esta actualización es móvil')
p('La propuesta aprovecha la cámara, los sensores biométricos, el GPS, la conexión NFC, el motor de vibración y el acelerómetro del teléfono. Estos recursos permiten que Nexa acompañe al usuario durante situaciones cotidianas: escanear un producto, confirmar una compra o consultar una entrega mientras se desplaza.')
p('Su carácter móvil se justifica por la combinación de portabilidad, ubicación e interacción directa con el entorno. Para ofrecer esta experiencia, será necesario integrar las funciones con el sistema operativo y comprobar su disponibilidad en cada dispositivo.')
h('Arquitectura de la aplicación')
p('Para esta actualización se propone utilizar una arquitectura cliente servidor con una aplicación móvil organizada por capas y un backend modular. Esta estructura permitirá integrar las funciones del smartphone y desarrollar cada etapa de forma independiente.')
sub('Capa de presentación')
p('Contendrá las pantallas de inicio de sesión, catálogo, carrito, pagos y seguimiento de pedidos. Mostrará la información y las respuestas visuales de cada interacción, incluido el mapa asociado a InteractiveMap.js.')
sub('Capa de lógica de la aplicación')
p('Organizará los procesos de acceso, búsqueda, compra, vinculación de dispositivos y consulta de pedidos. Coordinará las acciones de las pantallas con los servicios del servidor y las funciones del teléfono.')
sub('Capa de integración con el dispositivo')
p('Conectará Nexa con la autenticación biométrica, la cámara, el GPS, NFC, las notificaciones y los sensores de movimiento y vibración. Gestionará los permisos y comprobará la compatibilidad de cada función, ofreciendo alternativas cuando no esté disponible.')
sub('Backend modular y almacenamiento')
p('El servidor se organizará inicialmente como un monolito modular, es decir, una sola aplicación con módulos separados para usuarios, catálogo, pedidos, pagos y notificaciones. Estos módulos compartirán una base de datos según las necesidades del sistema.')
p('La aplicación móvil se comunicará con el backend mediante una API protegida con HTTPS. El seguimiento del repartidor incorporará un canal de actualización en tiempo real. Las operaciones sensibles, como validar precios, autorizar pedidos y confirmar pagos, se verificarán en el servidor.')
sub('Justificación de la arquitectura')
p('Esta arquitectura permitirá incorporar las mejoras por etapas sin mezclar la interfaz con el acceso al hardware ni con las reglas del negocio. El backend modular facilitará el mantenimiento inicial, mientras que la separación por capas permitirá adaptar las integraciones móviles a cada plataforma.')
p('Su elección responde al carácter móvil de Nexa: el teléfono gestionará la interacción con el usuario y sus sensores, y el servidor mantendrá la información compartida y la validación de las operaciones.')

stages=[
('Etapa 1 Acceso y confirmación mediante biometría',
'Se incorporará autenticación mediante reconocimiento facial, como Face ID, o huella dactilar para facilitar el inicio de sesión y la confirmación de compras. La biometría sustituirá la captura habitual del PIN cuando el dispositivo sea compatible y el usuario active esta opción.',
'El teléfono permite verificar la identidad mediante sus sensores y mecanismos de seguridad, reduciendo la necesidad de escribir credenciales en una pantalla pequeña.',
'Se utilizará la validación del sistema operativo, sin almacenar imágenes faciales ni huellas en Nexa. Se conservará un método alternativo de acceso y recuperación. La confirmación biométrica no reemplazará las validaciones necesarias para procesar el pago.',
'Un acceso más ágil y una confirmación de compra con menos pasos. La etapa se validará comprobando tanto la autenticación como el método alternativo.'),
('Etapa 2 Escaneo con la cámara',
'Se habilitará el escaneo de códigos QR para vincular dispositivos y de códigos de barras para localizar productos dentro del catálogo. El usuario abrirá el lector desde Nexa y apuntará la cámara al código.',
'La cámara convierte al teléfono en un lector portátil que conecta productos y dispositivos físicos con las funciones de la aplicación.',
'Se solicitará permiso al utilizar la cámara y se comprobará que el código corresponda a un producto o proceso válido. Se mantendrán la búsqueda manual y la introducción de códigos cuando el escaneo no sea posible.',
'Menos escritura y errores al buscar productos o vincular dispositivos. Se probarán códigos válidos, desconocidos y difíciles de leer.'),
('Etapa 3 Direcciones y seguimiento con GPS',
'La ubicación del teléfono se utilizará para sugerir una dirección de envío, que el usuario podrá corregir y confirmar. Además, se propone incorporar en InteractiveMap.js la visualización de la ubicación del repartidor durante una entrega activa.',
'El GPS permite adaptar la experiencia al lugar donde se encuentra el usuario y consultar el recorrido de una entrega desde cualquier ubicación con conexión.',
'Las coordenadas deberán convertirse en una dirección mediante un servicio de mapas; será necesario confirmar datos como número exterior o interior. El seguimiento dependerá de las ubicaciones enviadas por el repartidor y del servicio que las comunique a Nexa. El GPS del comprador, por sí solo, no permite rastrearlo.',
'Una captura de dirección más sencilla y mayor visibilidad de la entrega. El mapa deberá indicar cuándo se actualizó la posición y comunicar si no hay datos recientes.'),
('Etapa 4 Notificaciones push del pedido',
'Se integrarán notificaciones nativas para informar cambios como pedido confirmado, preparación, salida a reparto y entrega. Al tocar una alerta, el usuario accederá al pedido correspondiente.',
'Las notificaciones permiten recibir novedades sin mantener Nexa abierta y consultar la información desde el teléfono en el momento en que resulte relevante.',
'Se solicitará autorización para enviar alertas y se ofrecerán preferencias para distinguir avisos del pedido y promociones. El historial dentro de Nexa permitirá consultar el estado aunque una notificación no llegue o esté desactivada.',
'Menor necesidad de revisar manualmente los pedidos. Se validará que las alertas correspondan a cambios reales y abran el pedido correcto.'),
('Etapa 5 Funciones por proximidad con NFC',
'Se incorporará la vinculación por proximidad con dispositivos o etiquetas compatibles. Los pagos sin contacto se contemplarán mediante una integración de pago admitida por la plataforma y el proveedor seleccionado.',
'NFC permite iniciar determinadas acciones acercando el teléfono a un elemento compatible, lo que facilita interacciones presenciales.',
'Se verificará la compatibilidad del equipo y del sistema operativo. La disponibilidad de NFC no garantiza por sí sola que Nexa pueda realizar pagos. Se ofrecerán alternativas, como QR para vinculación y los métodos de pago disponibles en la aplicación.',
'Menos pasos en interacciones por proximidad. Primero se validará la vinculación; los pagos se habilitarán una vez comprobada la integración completa.'),
('Etapa 6 Respuestas táctiles y promociones por movimiento',
'Se añadirá feedback háptico, es decir, pequeñas respuestas de vibración, para acompañar acciones como agregar un producto al carrito. También se propone utilizar el acelerómetro para detectar un gesto de agitación que permita descubrir o activar promociones disponibles.',
'El motor de vibración ofrece una respuesta que el usuario puede sentir, mientras que el acelerómetro permite reconocer el movimiento del teléfono.',
'Ambas funciones podrán desactivarse. La vibración se acompañará de una confirmación visual y las promociones tendrán también un botón de acceso. Agitar el dispositivo no deberá realizar compras ni modificar el carrito de forma involuntaria.',
'Interacciones más claras y una forma opcional de descubrir promociones. Se probará que los movimientos cotidianos no provoquen activaciones accidentales.')]
phase_titles = ['Fase 1 Acceso seguro y búsqueda de productos', 'Fase 2 Entrega y comunicación en tiempo real', 'Fase 3 Proximidad e interacción con el dispositivo']
phase_objectives = [
    'Esta fase reúne la autenticación biométrica y el escaneo con la cámara para facilitar el acceso y la búsqueda. Se avanzará cuando ambas funciones y sus alternativas manuales funcionen correctamente.',
    'Esta fase reúne la ubicación GPS, el seguimiento en InteractiveMap.js y las notificaciones push. Se avanzará cuando el estado del pedido, el mapa y las alertas sean coherentes, incluso ante interrupciones de conexión.',
    'Esta fase incorpora NFC, respuestas hápticas y promociones mediante el acelerómetro. Su cierre requiere validar compatibilidad, alternativas accesibles y ausencia de activaciones involuntarias. Los pagos NFC dependerán de una integración compatible.'
]
for index,(title,description,reason,conditions,result) in enumerate(stages):
    if index % 2 == 0:
        h(phase_titles[index // 2])
        p(phase_objectives[index // 2])
    sub(title.split(' ', 2)[2])
    p(description)
    label('Justificación móvil.',reason)
    label('Condiciones de implementación.',conditions)
    label('Resultado esperado.',result)
h('Criterios para avanzar entre etapas')
p('Cada etapa deberá comprobarse en dispositivos compatibles, con permisos aceptados y rechazados, y ante interrupciones de conexión cuando corresponda. Las funciones esenciales conservarán alternativas para usuarios cuyos teléfonos no dispongan de determinados sensores.')
p('El orden propuesto prioriza el acceso y la búsqueda, continúa con la entrega y la comunicación del pedido, y deja para después las integraciones por proximidad y las interacciones complementarias. Así, cada incorporación puede evaluarse antes de ampliar el alcance.')
h('Conclusión')
p('La actualización dará a Nexa una experiencia móvil centrada en tareas concretas: acceder con biometría, escanear productos, completar direcciones, seguir entregas, recibir alertas e interactuar mediante proximidad y movimiento.')
p('El valor de estas funciones dependerá de que reduzcan pasos y faciliten el uso cotidiano. La implementación por etapas permitirá comprobar ese beneficio y mantener la aplicación accesible en teléfonos con distintas capacidades.')
doc.core_properties.title='Actualización móvil de Nexa por etapas'
doc.core_properties.subject='Arquitectura propuesta y etapas de implementación'
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.text = run.text.replace('en seis etapas', 'en tres fases').replace('por etapas', 'en tres fases').replace('cada etapa', 'cada fase').replace('Cada etapa', 'Cada fase').replace('La etapa', 'La función').replace('entre etapas', 'entre fases')
doc.core_properties.title='Actualización móvil de Nexa en tres fases'
doc.core_properties.subject='Arquitectura propuesta y tres fases de implementación'
out=Path(__file__).parent/'Actualizacion_movil_Nexa_tres_fases.docx'
doc.save(out)
check=Document(out)
assert len([q for q in check.paragraphs if q.text.startswith('Fase ')])==3
assert any(q.text=='Arquitectura de la aplicación' for q in check.paragraphs)
print(out.resolve())
