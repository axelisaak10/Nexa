from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output" / "word" / "Reporte_Proyecto_Integrador_CICD.docx"
EVID = ROOT / "evidencias"
TMP = ROOT / "tmp" / "report"
OUT.parent.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)
NAVY, TEAL, PALE, GRAY = "0D2B45", "007F73", "EAF3F6", "52606D"


def pil_font(size, bold=False):
    names = ["aptosbd.ttf", "arialbd.ttf"] if bold else ["aptos.ttf", "arial.ttf"]
    for name in names:
        candidate = Path("C:/Windows/Fonts") / name
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def architecture_image():
    path = TMP / "arquitectura-cicd.png"
    im = Image.new("RGB", (1600, 720), "white")
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((35, 35, 1565, 685), 28, fill="#F5F8FA", outline="#C8D6DF", width=3)
    d.text((70, 65), "Arquitectura del pipeline CI/CD", fill="#0D2B45", font=pil_font(42, True))
    boxes = [
        (80, 210, 330, 460, "GitHub", "Código fuente\ny eventos push/PR"),
        (410, 210, 690, 460, "GitHub Actions", "25 pruebas\ncobertura > 70 %"),
        (770, 210, 1040, 460, "Docker Hub", "Etiquetas latest\ny SHA del commit"),
        (1120, 210, 1500, 460, "AWS EC2", "Docker + Nginx\nblue/green · puerto 80"),
    ]
    for x1, y1, x2, y2, label, body in boxes:
        d.rounded_rectangle((x1, y1, x2, y2), 24, fill="white", outline="#007F73", width=5)
        lf, bf = pil_font(28, True), pil_font(21)
        lw = d.textlength(label, font=lf)
        d.text(((x1 + x2 - lw) / 2, y1 + 55), label, fill="#0D2B45", font=lf)
        bb = d.multiline_textbbox((0, 0), body, font=bf, spacing=12, align="center")
        d.multiline_text(((x1 + x2 - (bb[2] - bb[0])) / 2, y1 + 125), body, fill="#52606D", font=bf, spacing=12, align="center")
    for start, end in [((330, 335), (410, 335)), ((690, 335), (770, 335)), ((1040, 335), (1120, 335))]:
        d.line((*start, *end), fill="#F28C28", width=10)
        ex, ey = end
        d.polygon([(ex, ey), (ex - 24, ey - 16), (ex - 24, ey + 16)], fill="#F28C28")
    d.rounded_rectangle((410, 535, 1500, 625), 20, fill="#0D2B45")
    footer = "Secretos cifrados · verificación /api/health · activación sin tiempo de inactividad"
    d.text(((1910 - d.textlength(footer, font=pil_font(21))) / 2, 565), footer, fill="white", font=pil_font(21))
    im.save(path)
    return path


def crop(name, box, output_name):
    target = TMP / output_name
    with Image.open(EVID / name) as im:
        im.crop(box).save(target)
    return target


ARCH = architecture_image()
GH = crop("19-github-actions-cicd-success.png", (0, 0, 1440, 760), "github-actions-resumen.png")
DH = crop("20-dockerhub-tags.png", (0, 0, 1440, 1050), "dockerhub-tags-resumen.png")
AWS = crop("21-aws-api-publica.png", (0, 0, 1440, 1030), "aws-api-publica-resumen.png")

doc = Document()
section = doc.sections[0]
section.top_margin, section.bottom_margin = Inches(.72), Inches(.7)
section.left_margin, section.right_margin = Inches(.82), Inches(.82)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Aptos", Pt(10.5)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
    style = doc.styles[style_name]
    style.font.name = "Aptos Display"
    style.font.color.rgb = RGBColor.from_string(NAVY)
    style.paragraph_format.keep_with_next = True
doc.styles["Heading 1"].font.size = Pt(17)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 2"].font.size = Pt(13)
doc.styles["Heading 2"].font.bold = True


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, value=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar") or OxmlElement("w:tcMar")
    if tc_mar.getparent() is None:
        tc_pr.append(tc_mar)
    for key in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{key}")) or OxmlElement(f"w:{key}")
        if node.getparent() is None:
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def report_table(headers, rows, widths=None, size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style, table.autofit = "Table Grid", False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell, NAVY)
        cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb, run.font.bold, run.font.size = RGBColor(255, 255, 255), True, Pt(size)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            cell_margins(cells[i])
            if row_index % 2:
                shade(cells[i], PALE)
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(size)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    if level == 1:
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("val", "single"), ("sz", "12"), ("space", "5"), ("color", TEAL)):
            bottom.set(qn(f"w:{key}"), value)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)
    return p


def text(value):
    return doc.add_paragraph(value)


def bullet(value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Inches(.22), Inches(-.15)
    p.add_run(value)
    return p


def figure(path, number, caption, width=6.55):
    p = doc.add_paragraph()
    p.alignment, p.paragraph_format.keep_with_next = WD_ALIGN_PARAGRAPH.CENTER, True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(f"Figura {number}. {caption}")
    run.italic, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor.from_string(GRAY)


def code(lines):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, "102A38")
    cell_margins(cell, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(lines):
        run = p.add_run(line + ("\n" if i < len(lines) - 1 else ""))
        run.font.name, run.font.size, run.font.color.rgb = "Cascadia Mono", Pt(8.5), RGBColor(235, 245, 247)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page():
    doc.add_page_break()


# Portada
cover = doc.add_table(rows=1, cols=1)
cell = cover.cell(0, 0)
shade(cell, NAVY)
cell_margins(cell, 220)
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("UNIVERSIDAD TECNOLÓGICA DE QUERÉTARO")
r.bold, r.font.size, r.font.color.rgb = True, Pt(18), RGBColor(255, 255, 255)
p = cell.add_paragraph("Unidad I · Introducción a DevOps")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in p.runs:
    r.font.size, r.font.color.rgb = Pt(12), RGBColor.from_string("9DE5DB")
for _ in range(3): doc.add_paragraph()
p = doc.add_paragraph("PROYECTO INTEGRADOR")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.bold, p.runs[0].font.size, p.runs[0].font.color.rgb = True, Pt(12), RGBColor.from_string(TEAL)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Pipeline CI/CD automatizado\npara API REST")
r.bold, r.font.size, r.font.color.rgb = True, Pt(26), RGBColor.from_string(NAVY)
p = doc.add_paragraph("Docker · GitHub Actions · Docker Hub · AWS EC2")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(14), RGBColor.from_string(GRAY)
for _ in range(3): doc.add_paragraph()
report_table(["DATOS DE ENTREGA", "INFORMACIÓN"], [
    ["Alumno", "Axel Rodríguez"], ["Asignatura", "Introducción a DevOps"],
    ["Docente", "________________________________"], ["Grupo", "________________________________"],
    ["Fecha", "8 de octubre de 2026"], ["Repositorio", "github.com/axelisaak10/Nexa"],
], [2.0, 4.8], 9.5)
p = doc.add_paragraph("Santiago de Querétaro, Querétaro")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
page()

# Índices
heading("Índice", 1)
report_table(["Núm.", "Sección", "Pág."], [
    ("1", "Introducción", "3"), ("2", "Objetivo y alcance", "4"),
    ("3", "Arquitectura de la solución", "4"), ("4", "Desarrollo y calidad del backend", "5"),
    ("5", "Pruebas automatizadas y cobertura", "6"), ("6", "Contenedorización y Docker Hub", "7"),
    ("7", "Integración continua con GitHub Actions", "8"), ("8", "Despliegue continuo en AWS EC2", "9"),
    ("9", "Resultados y análisis", "10"), ("10", "Conclusiones", "11"),
    ("11", "Bibliografía", "12"), ("12", "Anexos de evidencias", "13"),
], [.7, 5.4, .7], 10)
heading("Índice de figuras", 2)
report_table(["Figura", "Descripción", "Pág."], [
    ("1", "Arquitectura del pipeline CI/CD", "4"), ("2", "Pruebas automatizadas y cobertura", "13"),
    ("3", "Colección ejecutada en Postman", "14"), ("4", "Respuesta JSON del endpoint de salud", "15"),
    ("5", "Pipeline completo en GitHub Actions", "16"), ("6", "Etiquetas de la imagen en Docker Hub", "17"),
    ("7", "Aplicación pública desplegada en AWS EC2", "18"), ("8", "Instancia EC2 en ejecución", "19"),
], [.8, 5.3, .7], 9.5)
page()

# Introducción
heading("1. Introducción", 1)
for paragraph in [
    "La integración continua y el despliegue continuo constituyen prácticas centrales de DevOps porque convierten los cambios de software en versiones verificadas, reproducibles y disponibles. En un proceso manual, el equipo debe recordar cómo instalar dependencias, ejecutar pruebas, construir una imagen, etiquetarla, publicarla y sustituir la aplicación que se encuentra en producción. Cada paso manual agrega variación y aumenta el riesgo de omitir una validación, utilizar una versión incorrecta o exponer información sensible. Este proyecto transforma ese proceso en un pipeline automatizado y auditable.",
    "La solución parte de una API REST implementada con Node.js y SQLite. El dominio utiliza categorías y productos relacionados mediante claves foráneas y ofrece doce operaciones HTTP con los métodos GET, POST, PUT y DELETE. Todas las respuestas mantienen el contrato JSON formado por statusCode y data. La API también incorpora validación de entrada, autenticación administrativa para escrituras, control de errores, respaldo de la base de datos, vaciado seguro y un endpoint de salud que actualmente responde con el mensaje hola.",
    "La calidad se comprueba antes de construir la imagen. La suite automatizada ejecuta veinticinco pruebas unitarias y de integración para los escenarios correctos y para fallos comunes: cuerpos inválidos, identificadores inexistentes, conflictos de integridad, credenciales incorrectas y operaciones destructivas sin confirmación. La cobertura resultante supera el umbral mínimo de 70 % en líneas, ramas y funciones. También se verifican el Dockerfile, las etiquetas de la imagen y la estructura del flujo de GitHub Actions.",
    "El pipeline se ejecuta con GitHub Actions en eventos push y pull_request dirigidos a la rama main. Primero instala las dependencias y ejecuta las pruebas con cobertura. Después de un push aprobado, construye una imagen Docker y publica dos referencias en Docker Hub: latest y el hash completo del commit. La etiqueta inmutable permite identificar exactamente qué código se desplegó, mientras que latest ofrece una referencia sencilla para la versión más reciente.",
    "El despliegue continuo conecta el workflow con una instancia Ubuntu de Amazon EC2 mediante una llave SSH almacenada en GitHub Secrets. La instancia descarga la imagen, inicia la versión nueva, consulta /api/health y activa el tráfico público en el puerto 80 únicamente cuando la comprobación es satisfactoria. Nginx permite alternar entre posiciones azul y verde, de modo que el contenedor anterior permanece disponible durante la verificación y se reduce el tiempo de interrupción.",
    "Este reporte documenta la arquitectura, las decisiones de implementación, las pruebas, la cobertura, la publicación en Docker Hub y el despliegue en AWS. Las evidencias se presentan como figuras numeradas y se relacionan con los resultados obtenidos. El objetivo final es demostrar que un cambio enviado al repositorio puede recorrer de forma segura todo el ciclo: validación, empaquetado, publicación y actualización automática del servidor.",
]: text(paragraph)
page()

# Objetivo y arquitectura
heading("2. Objetivo y alcance", 1)
text("Objetivo general. Implementar un pipeline CI/CD que valide una API REST, genere una imagen Docker, la publique en Docker Hub y despliegue automáticamente la versión aprobada en AWS EC2 por el puerto 80.")
for item in [
    "Mantener al menos seis endpoints funcionales; la versión entregada contiene doce operaciones.",
    "Ejecutar pruebas automáticas con un umbral mínimo de cobertura de 70 %.",
    "Construir una imagen optimizada con Dockerfile y un contexto reducido mediante .dockerignore.",
    "Publicar las etiquetas latest y el hash del commit en Docker Hub.",
    "Desplegar en EC2 con verificación de salud y estrategia azul/verde.",
    "Proteger llaves, tokens, direcciones y contraseñas mediante GitHub Secrets.",
]: bullet(item)
heading("3. Arquitectura de la solución", 1)
text("La Figura 1 resume el recorrido de un cambio desde el repositorio hasta el servicio público. Cada etapa depende del éxito de la anterior; por esta razón, una falla de pruebas impide publicar la imagen y una falla de salud impide activar la nueva versión.")
figure(ARCH, 1, "Arquitectura del pipeline CI/CD automatizado.", 5.9)
report_table(["Componente", "Responsabilidad", "Evidencia"], [
    ["GitHub", "Aloja el código fuente y recibe push/PR.", "Repositorio Nexa"],
    ["GitHub Actions", "Ejecuta pruebas, build, publicación y despliegue.", "Figura 5"],
    ["Docker Hub", "Conserva latest y la etiqueta SHA.", "Figura 6"],
    ["AWS EC2", "Ejecuta Docker y Nginx en Ubuntu.", "Figuras 7 y 8"],
    ["Postman", "Comprueba respuestas públicas y contratos JSON.", "Figuras 3 y 4"],
], [1.3, 3.9, 1.5], 7.8)
page()

# Backend
heading("4. Desarrollo y calidad del backend", 1)
text("La API administra un catálogo normalizado. Una categoría puede relacionarse con varios productos y cada producto conserva category_id como clave foránea. Los precios se almacenan como enteros en price_cents para evitar errores de redondeo. Las operaciones de escritura requieren el encabezado X-Admin-Token y los cuerpos se validan antes de modificar SQLite.")
report_table(["Método", "Ruta", "Operación"], [
    ["GET", "/api/health", "Estado del servicio; devuelve el mensaje hola"],
    ["GET", "/api/categories", "Listar categorías"], ["POST", "/api/categories", "Crear categoría"],
    ["PUT", "/api/categories/{id}", "Actualizar categoría"], ["DELETE", "/api/categories/{id}", "Eliminar categoría sin productos"],
    ["GET", "/api/products", "Listar productos"], ["POST", "/api/products", "Crear producto"],
    ["GET", "/api/products/{id}", "Consultar producto"], ["PUT", "/api/products/{id}", "Actualizar producto"],
    ["DELETE", "/api/products/{id}", "Eliminar producto"], ["POST", "/api/database/backup", "Crear respaldo consistente de SQLite"],
    ["DELETE", "/api/database", "Vaciar datos y conservar la estructura"],
], [.9, 2.6, 3.2], 7.3)
heading("4.1 Contrato de respuesta", 2)
text("El contrato uniforme facilita las pruebas automáticas y el consumo desde Postman. El endpoint GET /api/health produce la siguiente respuesta pública, visible nuevamente en la Figura 4:")
code(['{', '  "statusCode": 200,', '  "data": [{ "message": "hola" }]', '}'])
heading("4.2 Casos de error considerados", 2)
for item in [
    "400: cuerpo JSON incompleto, identificador inválido o confirmación incorrecta.",
    "401: ausencia del token administrativo o token inválido.",
    "404: categoría o producto inexistente.",
    "409: nombre duplicado o intento de eliminar una categoría con productos.",
    "500: error interno controlado con la misma envoltura JSON.",
]: bullet(item)
page()

# Pruebas
heading("5. Pruebas automatizadas y cobertura", 1)
text("El comando npm run test:coverage ejecuta la batería completa antes de que el pipeline permita construir la imagen. Las veinticinco pruebas aprobaron y no hubo casos omitidos. La Figura 2 presenta la salida del ejecutor; la Figura 3 complementa esta validación con veinte aserciones correctas realizadas sobre la colección de Postman.")
report_table(["Métrica", "Resultado", "Umbral", "Evaluación"], [
    ["Pruebas", "25 aprobadas / 0 fallidas", "0 fallas", "Cumplido"],
    ["Líneas", "93.36 %", "70 %", "Superado"], ["Ramas", "88.18 %", "70 %", "Superado"],
    ["Funciones", "93.55 %", "70 %", "Superado"],
], [1.5, 2.4, 1.3, 1.5], 9)
text("Las pruebas de integración levantan el servidor con una base aislada, realizan solicitudes HTTP reales y eliminan los datos temporales al finalizar. Esto evita que el resultado dependa de la base de producción. También se verifica el socket TCP de la práctica anterior y la persistencia compartida con la API HTTP.")
heading("5.1 Escenarios representativos", 2)
for item in [
    "Creación, consulta, actualización y eliminación de categorías y productos.",
    "Rechazo de nombres duplicados y precios inválidos.",
    "Integridad referencial entre categorías y productos.",
    "Autorización de operaciones de escritura mediante ADMIN_TOKEN.",
    "Respaldo coherente y vaciado confirmado de SQLite.",
    "Contrato {statusCode, data} en respuestas exitosas y fallidas.",
]: bullet(item)
page()

# Docker
heading("6. Contenedorización y Docker Hub", 1)
text("El Dockerfile emplea una imagen base de Node.js 22 sobre Debian slim, instala únicamente dependencias de producción, copia los archivos necesarios, crea el directorio persistente /data y ejecuta el proceso con el usuario node. El contenedor expone el puerto 80 y define una comprobación de salud sobre /api/health. El archivo .dockerignore excluye node_modules, archivos .env, registros, cobertura, reportes y datos locales.")
report_table(["Control", "Implementación", "Beneficio"], [
    ["Imagen base", "node:22.22.3-bookworm-slim", "Entorno reproducible"], ["Usuario", "node", "Menor privilegio"],
    ["Persistencia", "/data", "Conserva SQLite entre versiones"], ["Healthcheck", "GET /api/health", "Detecta una versión no funcional"],
    ["Contexto", ".dockerignore", "Reduce tamaño y exposición accidental"],
], [1.4, 2.5, 2.8], 9)
text("Después de aprobar las pruebas, el workflow publica 34axel/nexa-api:latest y 34axel/nexa-api:${{ github.sha }}. La Figura 6 muestra ambas referencias en Docker Hub. En la ejecución documentada, el hash inicia con aba4487a y permite relacionar la imagen con el commit exacto.")
code(["docker pull 34axel/nexa-api:latest", "docker pull 34axel/nexa-api:aba4487a19b8b37fd6def29e04b9532c8d2497ea"])
page()

# CI
heading("7. Integración continua con GitHub Actions", 1)
text("El archivo .github/workflows/main.yml se activa con push y pull_request hacia main. La ejecución registrada para el cambio feat: devolver mensaje JSON en health terminó con estado Success en 1 minuto y 6 segundos. La Figura 5 muestra el orden y la aprobación de los tres trabajos.")
report_table(["Trabajo", "Duración", "Resultado", "Acción principal"], [
    ["Pruebas y cobertura", "17 s", "Correcto", "npm run test:coverage"],
    ["Construir y publicar imagen", "33 s", "Correcto", "Buildx y push a Docker Hub"],
    ["Desplegar en AWS EC2", "9 s", "Correcto", "SSH, healthcheck y Nginx"],
], [2.0, 1.0, 1.2, 2.5], 8.7)
heading("7.1 Gestión segura de credenciales", 2)
text("La configuración sensible se almacena como GitHub Secrets y no aparece en el repositorio. El workflow consume las siguientes variables únicamente durante la ejecución:")
for item in [
    "DOCKERHUB_USERNAME y DOCKERHUB_TOKEN para publicar la imagen.",
    "EC2_HOST y EC2_USER para identificar el servidor de destino.",
    "EC2_SSH_KEY con el contenido privado del archivo .pem.",
    "ADMIN_TOKEN para proteger las rutas de escritura.",
    "DEPLOY_ENABLED como variable de control del despliegue.",
]: bullet(item)
text("La llave se escribe temporalmente con permisos restrictivos, SSH acepta la huella del servidor en el primer contacto y el archivo se elimina al terminar el trabajo.")
page()

# AWS
heading("8. Despliegue continuo en AWS EC2", 1)
text("La instancia Ubuntu ejecuta Docker y Nginx. El grupo de seguridad permite SSH por el puerto 22 para administración y HTTP por el puerto 80 para la API pública. Durante el despliegue se selecciona la posición inactiva, se descarga la etiqueta SHA, se inicia el contenedor nuevo y se consulta repetidamente /api/health.")
text("Solo después de obtener HTTP 200 se genera la configuración de Nginx y se recarga el servicio. Si la comprobación falla, el contenedor nuevo se elimina y la posición anterior conserva el tráfico. La Figura 7 muestra la aplicación pública en 107.21.88.69 y la Figura 8 confirma la instancia EC2 en estado En ejecución.")
report_table(["Paso", "Validación", "Resultado esperado"], [
    ["1. Pull", "Etiqueta SHA disponible", "Imagen exacta descargada"], ["2. Inicio", "Contenedor nuevo activo", "Puerto interno alterno"],
    ["3. Salud", "GET /api/health = 200", "Respuesta {message: hola}"], ["4. Activación", "nginx -t correcto", "Tráfico público al nuevo color"],
    ["5. Limpieza", "Versión anterior detenida", "Sin contenedores obsoletos"],
], [1.5, 2.5, 2.7], 8.8)
heading("8.1 Comprobación pública", 2)
code(["curl http://107.21.88.69/api/health", '{"statusCode":200,"data":[{"message":"hola"}]}'])
page()

# Resultados
heading("9. Resultados y análisis", 1)
text("La práctica alcanzó los resultados funcionales y de automatización solicitados. Las evidencias se concentran en las Figuras 2 a 8 y permiten seguir el recorrido desde las pruebas hasta la respuesta pública en EC2.")
report_table(["Requisito", "Evidencia cuantitativa", "Estado"], [
    ["API REST", "12 operaciones GET, POST, PUT y DELETE", "Cumplido"], ["Pruebas automáticas", "25 aprobadas; 0 fallidas", "Cumplido"],
    ["Cobertura", "93.36 % líneas; 88.18 % ramas; 93.55 % funciones", "Superado"], ["Postman", "20/20 aserciones aprobadas", "Cumplido"],
    ["Docker", "Dockerfile y .dockerignore funcionales", "Cumplido"], ["Docker Hub", "latest y SHA disponibles", "Cumplido"],
    ["GitHub Actions", "3/3 trabajos correctos", "Cumplido"], ["AWS EC2", "HTTP 200 en puerto 80", "Cumplido"],
    ["Health", "JSON con message: hola", "Cumplido"], ["Seguridad", "Secretos fuera del código público", "Cumplido"],
], [2.0, 3.7, 1.1], 8.2)
heading("9.1 Análisis", 2)
text("El umbral de cobertura no solo se alcanzó; se superó por más de dieciocho puntos en la métrica más baja. Esto ofrece margen para detectar regresiones, aunque no sustituye la revisión del comportamiento real. Por esa razón, la colección de Postman y la verificación pública complementan la suite interna.")
text("El uso conjunto de latest y SHA resuelve dos necesidades distintas: facilidad de uso y trazabilidad. EC2 despliega la referencia SHA, por lo que una investigación posterior puede identificar el commit exacto. La estrategia azul/verde agrega una barrera de protección: el tráfico cambia únicamente después de validar la nueva versión.")
text("La ejecución exitosa de GitHub Actions prueba la integración de las cuatro plataformas. El cambio del mensaje de salud llegó al repositorio, superó las pruebas, produjo una imagen nueva, fue publicado y finalmente respondió como JSON desde AWS. Este resultado satisface la demostración solicitada de realizar un git push y observar la actualización automática.")
page()

# Conclusiones
heading("10. Conclusiones", 1)
for paragraph in [
    "El proyecto integró desarrollo, pruebas, contenedorización y nube en un flujo único. La API no depende de una preparación manual para demostrar su calidad: cada cambio en main ejecuta veinticinco pruebas y aplica umbrales de cobertura antes de producir un artefacto. La cobertura obtenida, superior a 88 % en todas las métricas, confirma que se evaluaron tanto los caminos exitosos como errores frecuentes del consumidor.",
    "Docker Hub funciona como punto de intercambio entre integración y despliegue. Las etiquetas latest y SHA permiten disponer de la versión más reciente sin perder trazabilidad. El Dockerfile reduce privilegios, conserva SQLite en un volumen persistente y proporciona una comprobación de salud. Estas decisiones hacen que el mismo artefacto aprobado en GitHub Actions sea el que se ejecuta en EC2.",
    "El despliegue azul/verde mejora la disponibilidad porque separa la preparación de la activación. El nuevo contenedor recibe tráfico únicamente después de responder HTTP 200, mientras Nginx conserva la posibilidad de mantener la versión anterior ante una falla. La evidencia pública del mensaje hola demuestra que el cambio atravesó el pipeline completo y no quedó limitado al entorno local.",
    "La gestión de secretos completa la solución desde el punto de vista operativo. Los tokens de Docker Hub, la llave .pem, el host y el token administrativo permanecen fuera del código. Como mejora futura, conviene agregar HTTPS con un dominio, monitoreo de métricas y alertas, rotación periódica de credenciales y una base administrada cuando el volumen de información exceda las necesidades de SQLite.",
]: text(paragraph)
page()

# Bibliografía
heading("11. Bibliografía", 1)
text("Las fuentes se consultaron el 8 de octubre de 2026. Se priorizó documentación oficial para mantener precisión técnica.")
sources = [
    "Amazon Web Services. (2026). Get started with Amazon EC2. https://docs.aws.amazon.com/AWSEC2/latest/UserGuide/EC2_GetStarted.html",
    "Amazon Web Services. (2026). Security group rules. https://docs.aws.amazon.com/AWSEC2/latest/UserGuide/security-group-rules.html",
    "Docker. (2026). Dockerfile reference. https://docs.docker.com/reference/dockerfile/",
    "Docker. (2026). Build and push Docker images with GitHub Actions. https://docs.docker.com/build/ci/github-actions/",
    "Docker. (2026). Personal access tokens. https://docs.docker.com/security/access-tokens/",
    "GitHub. (2026). Understanding GitHub Actions. https://docs.github.com/actions/about-github-actions/understanding-github-actions",
    "GitHub. (2026). Using secrets in GitHub Actions. https://docs.github.com/actions/security-guides/using-secrets-in-github-actions",
    "Node.js. (2026). Test runner and collecting code coverage. https://nodejs.org/api/test.html",
    "Nginx. (2026). Beginner's guide: configuration and reload. https://nginx.org/en/docs/beginners_guide.html",
    "SQLite. (2026). SQLite backup API. https://www.sqlite.org/backup.html",
]
for index, source in enumerate(sources, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Inches(.28), Inches(-.28)
    p.add_run(f"{index}. {source}")
page()

# Anexos
heading("12. Anexos de evidencias", 1)
text("Las ocho figuras siguientes fueron mencionadas y analizadas en las secciones anteriores. En conjunto documentan las pruebas, el registro de imágenes, la automatización y la infraestructura en la nube.")
heading("Anexo A. Pruebas automatizadas", 2)
figure(EVID / "17-pruebas-unitarias-node.png", 2, "Veinticinco pruebas aprobadas y cobertura superior al 70 %.", 6.6)
page()
heading("Anexo B. Pruebas en Postman", 2)
figure(EVID / "16-postman-runner-resultados.png", 3, "Ejecución de la colección con veinte aserciones aprobadas y cero errores.", 6.55)
text("La ejecución recorre los endpoints en orden y valida el código HTTP y la envoltura JSON. El resultado confirma que la API puede consumirse desde una herramienta externa.")
page()
heading("Anexo C. Respuesta pública en Postman", 2)
figure(EVID / "14-postman-health-aws.png", 4, "Respuesta HTTP 200 del endpoint de salud consultado desde Postman.", 6.55)
text("Después de la actualización, el mismo endpoint conserva el contrato {statusCode, data} y entrega el mensaje hola. La respuesta exacta también se documenta en las secciones 4.1 y 8.1.")
page()
heading("Anexo D. Ejecución del pipeline", 2)
figure(GH, 5, "Tres trabajos correctos en GitHub Actions: pruebas, publicación y despliegue.", 6.55)
text("El resumen muestra el estado Success, el commit aba4487a y el flujo secuencial completado en 1 minuto y 6 segundos.")
page()
heading("Anexo E. Registro de contenedores", 2)
figure(DH, 6, "Docker Hub con las etiquetas del hash del commit y latest.", 6.55)
text("La etiqueta SHA distingue cada versión y evita depender únicamente de una referencia mutable.")
page()
heading("Anexo F. Aplicación desplegada", 2)
figure(AWS, 7, "Interfaz pública servida desde AWS EC2 por el puerto 80.", 6.55)
text("La tarjeta superior muestra hola y la tabla documenta las doce operaciones disponibles en la API.")
page()
heading("Anexo G. Infraestructura AWS", 2)
figure(EVID / "22-aws-instancia-ec2.png", 8, "Instancia EC2 Ubuntu en ejecución y con comprobaciones aprobadas.", 6.55)
text("La instancia aloja Docker y Nginx. El acceso público se concentra en HTTP 80 y la administración utiliza SSH 22.")

for section in doc.sections:
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("PROYECTO INTEGRADOR · DEVOPS")
    run.font.size, run.font.color.rgb = Pt(8), RGBColor.from_string(GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("UTEQ · Axel Rodríguez   |   Página ")
    run.font.size, run.font.color.rgb = Pt(8), RGBColor.from_string(GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

doc.core_properties.title = "Proyecto Integrador: Pipeline CI/CD automatizado para API REST"
doc.core_properties.subject = "Docker, GitHub Actions, Docker Hub y AWS EC2"
doc.core_properties.author = "Axel Rodríguez"
for paragraph in doc.paragraphs:
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:widowControl"))
doc.save(OUT)
print(OUT)
