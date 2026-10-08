from docx import Document
from pathlib import Path
p=Path('Reporte_Pruebas_Unitarias_API_Nexa.docx')
d=Document(p)
assert any(x.text == 'Reporte de Pruebas Unitarias de la API Nexa' for x in d.paragraphs)
assert len(d.tables) == 5
assert any('Matriz de los 12 endpoints' in x.text for x in d.paragraphs)
assert any('17 aprobadas, 0 fallidas' in x.text for x in d.paragraphs)
assert any('Evidencias visuales' in x.text for x in d.paragraphs)
assert len(d.inline_shapes) == 5
assert p.stat().st_size > 100000
print(f'OK: {len(d.paragraphs)} párrafos, {len(d.tables)} tablas, {len(d.inline_shapes)} imágenes, {p.stat().st_size} bytes')
