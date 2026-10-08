from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'Reporte_Pruebas_Unitarias_API_Nexa.docx'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.7)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10)
styles['Title'].font.name = 'Aptos Display'; styles['Title'].font.size = Pt(24); styles['Title'].font.bold = True; styles['Title'].font.color.rgb = RGBColor(0,0,0)
for name, size in [('Heading 1',16),('Heading 2',12)]:
    styles[name].font.name='Aptos Display'; styles[name].font.size=Pt(size); styles[name].font.bold=True; styles[name].font.color.rgb=RGBColor(0,0,0)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def margins(cell, top=90, start=100, bottom=90, end=100):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for k,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        e=OxmlElement('w:'+k); e.set(qn('w:w'),str(v)); e.set(qn('w:type'),'dxa'); mar.append(e)
def table(rows, widths=None):
    t=doc.add_table(rows=len(rows), cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            c=t.cell(i,j); c.text=str(val); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
            if widths: c.width=Inches(widths[j])
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
                for r in p.runs: r.font.size=Pt(8.5)
            if i==0:
                shade(c,'193B5A')
                for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
            elif i%2==0: shade(c,'EAF1F6')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def para(text='', bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.08
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold=True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p
def heading(text, level=1):
    p=doc.add_heading(text, level=level); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(5); return p

for _ in range(4): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('UNIDAD I   TEMA 1'); r.font.size=Pt(12); r.font.bold=True; r.font.color.rgb=RGBColor(26,93,135)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Reporte de Pruebas Unitarias de la API Nexa')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Introducción a DevOps'); r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=RGBColor(64,86,109)
doc.add_paragraph()
table([['Dato','Información'],['Actividad','Pruebas unitarias de endpoints previamente creados'],['Alumno','Axel Rodríguez'],['Tecnología','Node.js 22, node:test, assert y SQLite'],['Fecha de entrega','2 de octubre de 2026'],['Resultado','17 pruebas aprobadas de 17; 0 fallos']],[2.0,4.8])
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(28)
r=p.add_run('API desplegada en contenedor Docker y AWS EC2'); r.font.size=Pt(11); r.font.italic=True; r.font.color.rgb=RGBColor(64,86,109)
doc.add_page_break()
heading('Resumen ejecutivo')
para('Se implementó y ejecutó una suite automatizada sobre la API desarrollada en la práctica anterior. La prueba crea una base SQLite temporal y aislada, inicia el servidor en un puerto disponible y valida el contrato JSON {statusCode, data}. La ejecución cubrió lectura, creación, actualización, eliminación, respaldo, vaciado y comunicación TCP.')
para('Para cumplir el requisito de actualización se incorporaron dos operaciones PUT: una para categorías y otra para productos. La documentación OpenAPI se actualizó a 12 operaciones. La suite también reproduce errores comunes del consumidor: datos incompletos, identificadores inválidos, token incorrecto, JSON mal formado, tipo de contenido ausente, duplicados y relaciones inexistentes.')
heading('Conclusión')
para('El resultado fue satisfactorio: 17 de 17 pruebas aprobaron. La API mantiene la estructura uniforme de respuesta, protege las escrituras, valida los cuerpos con JSON Schema, conserva la integridad referencial y genera respaldos SQLite coherentes.')

doc.add_page_break(); heading('Objetivo y estrategia de prueba')
para('El objetivo fue comprobar que los endpoints responden correctamente ante solicitudes válidas y rechazan de forma controlada las entradas incorrectas. Se utilizó el módulo estándar node:test, equivalente funcional a Jest para esta práctica, junto con assert/strict.')
heading('Preparación del entorno',2); para('Cada ejecución crea una carpeta temporal, inicializa el esquema normalizado de categorías y productos, configura un token administrativo exclusivo de prueba y levanta servidores HTTP y TCP locales. Al finalizar se cierran los servidores y se elimina la base temporal. Así, las pruebas no alteran la base desplegada en EC2.')
heading('Comando de ejecución',2); p=doc.add_paragraph(); p.style='Intense Quote'; p.add_run('node --test --test-reporter spec test.mjs').font.name='Consolas'
heading('Criterios verificados',2); para('Cada respuesta HTTP debe incluir solamente statusCode y data; data siempre debe ser un arreglo. También se compara el código HTTP real con statusCode, se revisan los datos devueltos y se confirma la persistencia después de reiniciar el servidor.')
heading('Archivos principales',2); table([['Archivo','Función'],['server.mjs','Servidor HTTP, Socket TCP y acceso a SQLite'],['contracts.mjs','JSON Schema, validador y documentación OpenAPI'],['test.mjs','Suite automatizada de 17 pruebas'],['schema.sql','Tablas normalizadas, restricciones e índice']],[2.0,4.8])

doc.add_page_break(); heading('Matriz de los 12 endpoints probados')
para('La matriz incluye operaciones GET, POST, PUT y DELETE. PUT representa la actualización solicitada en la actividad.')
endpoints=[['Núm.','Método','Ruta','Comportamiento esperado','HTTP'],['1','GET','/api/health','Estado del servicio y SQLite','200'],['2','GET','/api/categories','Lista categorías','200'],['3','POST','/api/categories','Crea categoría','201'],['4','PUT','/api/categories/{id}','Actualiza categoría','200'],['5','DELETE','/api/categories/{id}','Elimina categoría sin productos','200'],['6','GET','/api/products','Lista productos','200'],['7','GET','/api/products/{id}','Consulta producto por ID','200'],['8','POST','/api/products','Crea producto relacionado','201'],['9','PUT','/api/products/{id}','Actualiza producto','200'],['10','DELETE','/api/products/{id}','Elimina producto','200'],['11','POST','/api/database/backup','Respalda SQLite','200'],['12','DELETE','/api/database','Vacía datos y conserva estructura','200']]
table(endpoints,[.45,.65,1.75,3.25,.55])

doc.add_page_break(); heading('Escenarios de fallo considerados')
para('Los casos siguientes representan errores previsibles del usuario. Todos devuelven una respuesta JSON controlada y un código HTTP adecuado.')
errors=[['Entrada o acción','Endpoint','Código','Validación'],['Body vacío, arreglo o null','POST categoría','400','Incumple JSON Schema'],['Nombre vacío o propiedad extra','POST categoría','400','Entrada inválida'],['Precio negativo, decimal o texto','POST producto','400','Tipo o rango inválido'],['Categoría inexistente','POST producto','409','Integridad referencial'],['Categoría duplicada','POST categoría','409','Restricción UNIQUE'],['Eliminar categoría con productos','DELETE categoría','409','Protección de relación'],['Token incorrecto o ausente','POST PUT DELETE','401','Acceso rechazado'],['Confirmación distinta de VACIAR','DELETE base','400','Evita borrado accidental'],['ID inexistente','GET o PUT producto','404','Registro no encontrado'],['ID no numérico','GET producto','400','Identificador inválido'],['JSON mal formado','POST categoría','400','JSON inválido'],['Sin application/json','POST categoría','415','Tipo requerido'],['Body mayor a 16 KB','POST categoría','413','Cuerpo demasiado grande'],['Ruta o método inexistente','GET o PATCH','404 o 405','Error controlado']]
table(errors,[2.0,1.35,.65,2.75])

doc.add_page_break(); heading('Resultados obtenidos')
table([['Métrica','Resultado'],['Pruebas ejecutadas','17'],['Aprobadas','17'],['Fallidas','0'],['Canceladas','0'],['Omitidas','0'],['Cobertura HTTP','12 endpoints'],['Prueba adicional','Socket TCP 6061 con insert y get']],[3.2,3.6])
heading('Evidencia resumida de la ejecución',2)
for text in ['GET health y documentación OpenAPI','GET, POST, PUT y DELETE de categorías','GET, POST, PUT y DELETE de productos','Respaldo SQLite e integrity_check = ok','Persistencia después de reiniciar','Vaciado conservando tablas y respaldos','Errores de esquema, autorización e integridad','Socket TCP 6061 insert y get']:
    p=doc.add_paragraph(style='List Bullet'); p.add_run(text)
para('Resultado: 17 aprobadas, 0 fallidas.', 'Resultado:')
heading('Interpretación',2); para('La aprobación total demuestra que las funciones principales cumplen sus contratos y que las entradas inválidas no producen respuestas ambiguas ni corrupción de datos. El respaldo pasó PRAGMA integrity_check y conservó el registro creado antes de la copia.')
heading('Limitaciones',2); para('Estas pruebas validan comportamiento funcional e integración local con SQLite. No sustituyen pruebas de carga, seguridad ofensiva ni disponibilidad prolongada de EC2. El despliegue debe ejecutar la misma versión del código para disponer de los nuevos endpoints PUT.')
heading('Recomendaciones',2); para('Ejecutar npm test antes de construir cada imagen Docker, agregar la prueba al flujo de integración continua y desplegar una nueva imagen en EC2 después de cualquier cambio en server.mjs o contracts.mjs.')

doc.add_page_break(); heading('Evidencias visuales')
para('Las capturas siguientes documentan la ejecución automatizada, la disponibilidad del servicio en AWS, las respuestas HTTP y la comunicación mediante Socket TCP 6061.')
evidencias=[
    ('17-pruebas-unitarias-node.png','Figura 1. Ejecución de la suite automatizada: 17 pruebas aprobadas y 0 fallidas.'),
    ('14-postman-health-aws.png','Figura 2. Endpoint de salud ejecutado desde Postman con respuesta HTTP 200 y SQLite disponible.'),
    ('16-postman-runner-resultados.png','Figura 3. Resultado del corredor de Postman sobre los endpoints desplegados en EC2.'),
    ('socket-6061-local.png','Figura 4. Prueba de inserción y consulta mediante Socket TCP en el puerto 6061.'),
    ('13-panel-aws-ec2.png','Figura 5. Interfaz de la aplicación web y catálogo de endpoints de la práctica.'),
]
for idx,(name,caption) in enumerate(evidencias):
    if idx: doc.add_page_break()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT/'evidencias'/name), width=Inches(6.8))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.keep_with_next=True
    r=c.add_run(caption); r.font.size=Pt(9); r.font.italic=True; r.font.color.rgb=RGBColor(64,86,109)

for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Unidad I - Tema 1 - Introducción a DevOps    |    ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)

doc.save(OUT)
print(OUT)
