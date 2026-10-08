from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
EVID = ROOT / "evidencias"
OUT = ROOT / "Reporte_Mejora_Socket_TCP_6061_EC2.docx"
NAVY, TEAL, PALE, GRAY = "17324D", "0B7A75", "F4F7F9", "506070"

def font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill); pr.append(node)

def margins(cell, value=110):
    pr = cell._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar"); pr.append(mar)
    for side in ("top", "start", "bottom", "end"):
        n = OxmlElement(f"w:{side}"); n.set(qn("w:w"), str(value)); n.set(qn("w:type"), "dxa"); mar.append(n)

def make_terminal(txt_path, png_path, title):
    lines = Path(txt_path).read_text(encoding="utf-8-sig").splitlines()
    fpath = Path("C:/Windows/Fonts/consola.ttf")
    font_m = ImageFont.truetype(str(fpath), 18) if fpath.exists() else ImageFont.load_default()
    font_b = ImageFont.truetype(str(fpath), 20) if fpath.exists() else ImageFont.load_default()
    w, line_h = 1500, 30
    h = 85 + line_h * len(lines) + 35
    im = Image.new("RGB", (w, h), "#10161d")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, w, 58), fill="#1c2631")
    for x, c in [(22,"#ff5f57"),(52,"#febc2e"),(82,"#28c840")]: d.ellipse((x,19,x+18,37), fill=c)
    d.text((120,16), title, fill="#e7edf3", font=font_b)
    y = 76
    for line in lines:
        color = "#80cbc4" if line.startswith("Resultado") else ("#8bd5ff" if line.startswith(">") else "#e7edf3")
        d.text((28,y), line, fill=color, font=font_m); y += line_h
    im.save(png_path)

make_terminal(EVID / "socket-6061-local.txt", EVID / "socket-6061-local.png", "Prueba local Socket TCP 6061")
make_terminal(EVID / "socket-6061-ec2.txt", EVID / "socket-6061-ec2.png", "Prueba pública AWS EC2 3.15.3.74:6061")

doc = Document(); sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65); sec.left_margin = sec.right_margin = Inches(.75)
normal = doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(10.4)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.08
for name,size,color in [("Title",28,NAVY),("Heading 1",19,NAVY),("Heading 2",13,TEAL)]:
    s=doc.styles[name]; s.font.name="Aptos Display"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(8); s.paragraph_format.space_after=Pt(7); s.paragraph_format.keep_with_next=True

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Nexa API  |  Mejora Socket TCP 6061  |  AWS EC2"), size=8, color=GRAY)

def p(text="", bold_prefix=None, align=None, style=None):
    q=doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        a=q.add_run(bold_prefix); a.bold=True; q.add_run(text[len(bold_prefix):])
    else: q.add_run(text)
    if align is not None: q.alignment=align
    return q

def bullet(text):
    q=doc.add_paragraph(style="List Bullet"); q.paragraph_format.space_after=Pt(3); q.add_run(text)

def code(text):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    c=t.cell(0,0); shade(c,"EEF2F5"); margins(c,140)
    q=c.paragraphs[0]; q.paragraph_format.space_after=Pt(0)
    for i,line in enumerate(text.splitlines()):
        if i: q.add_run().add_break()
        font(q.add_run(line), name="Consolas", size=8.5, color="263746")
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.style="Table Grid"
    for i,h in enumerate(headers):
        c=t.cell(0,i); c.text=h; c.width=Inches(widths[i]); shade(c,NAVY); margins(c)
        for r in c.paragraphs[0].runs: font(r,size=8.5,bold=True,color="FFFFFF")
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].width=Inches(widths[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cells[i])
            if ri%2: shade(cells[i],PALE)
            for q in cells[i].paragraphs:
                q.paragraph_format.space_after=Pt(0)
                for r in q.runs: font(r,size=8.3,color="263746")
    doc.add_paragraph()

def image(path, caption, width=7.0):
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.add_run().add_picture(str(path), width=Inches(width))
    c=doc.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in c.runs: font(r,size=8.5,color=GRAY)

def page(): doc.add_page_break()

# Portada
for _ in range(4): doc.add_paragraph()
p("Reporte de mejora Socket TCP 6061 en AWS EC2", align=WD_ALIGN_PARAGRAPH.CENTER, style="Title")
p("Práctica Ejecutar Web App en Contenedor", align=WD_ALIGN_PARAGRAPH.CENTER)
p("Unidad I  Tema 1  Introducción a DevOps", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
table(["Elemento","Resultado"],[("Aplicación","Nexa API con SQLite"),("Infraestructura","Contenedor Docker en EC2"),("Mejora","Socket TCP en puerto 6061"),("Fecha de verificación","25 de septiembre de 2026")],[2.2,4.5])
p("Resultado principal: la API HTTP continúa disponible en el puerto 8080 y el mismo contenedor atiende el protocolo TCP solicitado en el puerto 6061.", align=WD_ALIGN_PARAGRAPH.CENTER)

page(); doc.add_heading("1 Objetivo y alcance",1)
p("La mejora agrega comunicación directa mediante sockets TCP sin sustituir la API REST existente. El servidor HTTP y el servidor TCP se ejecutan en el mismo proceso de Node.js, comparten el mismo archivo SQLite y se publican desde un solo contenedor Docker en la instancia EC2.")
doc.add_heading("Requerimientos cubiertos",2)
table(["Requerimiento","Implementación","Estado"],[("Socket TCP","Servidor basado en node:net","Cumplido"),("Puerto 6061","Escucha y publicación 6061:6061","Cumplido"),("Formato insert","{insert:<element>}","Cumplido"),("Formato get","{get:<element>}","Cumplido"),("Mismo contenedor","HTTP y TCP en webapp:socket-6061","Cumplido"),("Persistencia","Volumen webapp-data con SQLite","Cumplido"),("Evidencias PDF","Pruebas local y pública documentadas","Cumplido")],[2.0,3.7,1.0])
doc.add_heading("Resultado verificable",2)
bullet("HTTP: http://3.15.3.74:8080/api/health respondió statusCode 200 y database sqlite-ok.")
bullet("TCP: 3.15.3.74:6061 aceptó inserción y consulta con respuestas JSON.")
bullet("Docker: el contenedor publica 0.0.0.0:8080->80 y 0.0.0.0:6061->6061.")

page(); doc.add_heading("2 Diseño del protocolo TCP",1)
p("Cada solicitud se envía como una línea terminada en salto de línea. El contenido entre llaves identifica la operación y contiene el body JSON del elemento. La respuesta conserva la envoltura usada por la práctica anterior.")
doc.add_heading("Insertar elementos",2)
code('{insert:{"name":"Adaptador TCP","price_cents":6061,"category_id":1}}\n\n{"statusCode":201,"data":[{"id":1,"name":"Adaptador TCP",\n"price_cents":6061,"category_id":1}]}')
doc.add_heading("Obtener elementos",2)
code('{get:{"entity":"product","id":1}}\n\n{"statusCode":200,"data":[{"id":1,"name":"Adaptador TCP",\n"price_cents":6061,"category_id":1}]}')
p("También se puede consultar usando el body previamente insertado: {get:{\"name\":\"Adaptador TCP\",\"price_cents\":6061,\"category_id\":1}}.")
doc.add_heading("Validaciones",2)
bullet("Categorías: name de 1 a 80 caracteres.")
bullet("Productos: name, price_cents y category_id válidos conforme al JSON Schema existente.")
bullet("Límite de mensaje de 16 KiB y tiempo de espera de 15 segundos.")
bullet("Errores con la misma forma {statusCode, data} y códigos 400, 404, 409, 413 o 500.")

page(); doc.add_heading("3 Cambios realizados",1)
table(["Archivo","Cambio"],[("server.mjs","Servidor TCP, parser de insert/get, validación y consultas SQLite."),("test.mjs","Prueba integrada de inserción, consulta y persistencia compartida."),("Dockerfile","SOCKET_PORT=6061 y EXPOSE 80 6061."),("socket-client.mjs","Cliente de línea de comandos para pruebas TCP."),("run-ec2.sh","Publicación adicional -p 6061:6061."),("README.md","Uso, ejemplos, despliegue y notas de seguridad.")],[2.0,4.7])
doc.add_heading("Inicio de ambos servidores",2)
code("HTTP  http://0.0.0.0:80\nTCP   0.0.0.0:6061\nBase de datos compartida  /data/app.db")
doc.add_heading("Construcción y ejecución",2)
code("docker build -t webapp:socket-6061 .\n\ndocker run -d --name webapp-socket-container \\\n  --restart unless-stopped \\\n  -p 8080:80 -p 6061:6061 \\\n  -v webapp-data:/data webapp:socket-6061")

page(); doc.add_heading("4 Configuración en AWS EC2",1)
p("La mejora se desplegó en la instancia i-0ea7789720e6b8bd5, con nombre nexa-devops-ubuntu y dirección IPv4 pública 3.15.3.74.")
inst=Path("C:/Users/axeli/AppData/Local/Temp/codex-clipboard-417df30b-dad1-402e-86b9-32f1fe9abb30.png")
if inst.exists(): image(inst,"Evidencia de la instancia EC2 en ejecución.",6.9)
doc.add_heading("Regla de entrada",2)
table(["Dato","Valor"],[("Grupo de seguridad","sg-09a9ab244af268868"),("Regla","sgr-0463f11ec01e648f8"),("Tipo","TCP personalizado"),("Puerto","6061"),("Origen temporal para revisión","0.0.0.0/0"),("Descripción","Socket TCP 6061 practica DevOps")],[2.4,4.3])
p("Seguridad: para la revisión en clase el puerto está accesible públicamente. Como insert no incluye autenticación en el formato solicitado, conviene limitar el origen a la IP del docente o cerrar la regla después de la evaluación.", bold_prefix="Seguridad:")

page(); doc.add_heading("5 Evidencia de prueba local",1)
p("Antes del despliegue se ejecutó la prueba integrada del servidor y una secuencia real de inserción y consulta en el puerto 6061.")
image(EVID / "socket-6061-local.png","Figura 1. Inserción y consulta local mediante Socket TCP.",7.0)
doc.add_heading("Pruebas automatizadas",2)
p("La suite completa finalizó con 15 pruebas aprobadas y 0 fallos. La prueba TCP confirma que el dato insertado por socket queda disponible en la misma base SQLite usada por la API.")

page(); doc.add_heading("6 Evidencia de prueba pública en EC2",1)
p("La siguiente prueba se originó desde la computadora local y se conectó a la IP pública de la instancia. Por ello valida simultáneamente la regla de seguridad, la publicación Docker del puerto y el servidor TCP dentro del contenedor.")
image(EVID / "socket-6061-ec2.png","Figura 2. Prueba pública contra 3.15.3.74:6061.",7.0)
doc.add_heading("Interpretación",2)
bullet("insert devolvió statusCode 201 y creó el producto con identificador 5.")
bullet("get devolvió statusCode 200 y recuperó exactamente el elemento insertado.")
bullet("El valor price_cents=6061 demuestra que se transmitió el body completo.")

page(); doc.add_heading("7 Verificación final y conclusión",1)
table(["Componente","Evidencia","Resultado"],[("Servidor HTTP","GET /api/health","200 y sqlite-ok"),("Servidor TCP","Conexión pública a 3.15.3.74:6061","Operativo"),("Operación insert","Producto Adaptador TCP EC2","Creado"),("Operación get","Consulta por el body previo","Recuperado"),("Contenedor","Puertos 8080 y 6061 publicados","Operativo"),("Persistencia","Volumen webapp-data","Conservada")],[1.7,3.6,1.4])
p("La mejora quedó implementada y desplegada en AWS EC2. La aplicación mantiene los diez endpoints HTTP de la práctica anterior y agrega el protocolo Socket TCP solicitado en el mismo contenedor. Las pruebas confirman que insert y get funcionan sobre SQLite usando el formato requerido.")
doc.add_heading("Comandos para demostración en clase",2)
code("node socket-client.mjs 3.15.3.74 6061 \\\n'{insert:{\"name\":\"Demo\",\"price_cents\":6061,\"category_id\":7}}'\n\nnode socket-client.mjs 3.15.3.74 6061 \\\n'{get:{\"name\":\"Demo\",\"price_cents\":6061,\"category_id\":7}}'")
p("Nota: si la dirección IPv4 pública cambia al detener e iniciar la instancia, debe actualizarse el host usado por el cliente y, si corresponde, las evidencias.")

doc.core_properties.title="Reporte de mejora Socket TCP 6061 en AWS EC2"
doc.core_properties.subject="Práctica de DevOps con Docker, EC2, SQLite y Socket TCP"
doc.core_properties.author="Axel Isaac Rodriguez Rangel"
doc.save(OUT)
print(OUT)
